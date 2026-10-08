from pathlib import Path

import pytest
from openpyxl import load_workbook

from reporte import MONEDA, generar_reporte

DATA = Path(__file__).resolve().parent.parent / "data"


@pytest.fixture(scope="module")
def reporte(tmp_path_factory):
    ruta = tmp_path_factory.mktemp("salida") / "conciliacion.xlsx"
    generar_reporte(DATA / "extracto_banco.csv", DATA / "libro_bancos.xlsx", ruta)
    return load_workbook(ruta)


def test_tiene_todas_las_hojas(reporte):
    assert reporte.sheetnames == [
        "Resumen", "Conciliados", "Solo en banco", "Solo en registro", "Diferencias de monto", "Duplicados",
    ]


@pytest.mark.parametrize("hoja, filas", [
    ("Conciliados", 110),
    ("Solo en banco", 9),
    ("Solo en registro", 5),
    ("Diferencias de monto", 3),
    ("Duplicados", 2),
])
def test_cada_hoja_tiene_sus_movimientos(reporte, hoja, filas):
    assert reporte[hoja].max_row - 1 == filas


def test_los_montos_se_guardan_como_numeros(reporte):
    hoja = reporte["Solo en banco"]
    for celda in hoja["D"][1:]:
        assert isinstance(celda.value, (int, float))
        assert celda.number_format == MONEDA


def test_las_operaciones_conservan_los_ceros(reporte):
    operaciones = [c.value for c in reporte["Solo en banco"]["B"][1:]]
    assert all(isinstance(op, str) and len(op) == 8 for op in operaciones)
    assert any(op.startswith("0") for op in operaciones)


def test_el_resumen_calcula_con_formulas(reporte):
    resumen = reporte["Resumen"]
    assert resumen["B16"].value == "=SUM(B13:B15)"
    assert resumen["B23"].value == "=SUM(B18:B22)"
    assert resumen["B25"].value == "=B16-B23"


def test_las_hojas_tienen_filtros_y_encabezado_fijo(reporte):
    for nombre in reporte.sheetnames[1:]:
        assert reporte[nombre].freeze_panes == "A2"
        assert reporte[nombre].auto_filter.ref
