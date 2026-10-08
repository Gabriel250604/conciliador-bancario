import argparse
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from conciliacion import conciliar, cuadre
from lectura import leer_extracto, leer_libro

MONEDA = "#,##0.00;[Red]-#,##0.00"
FECHA = "dd/mm/yyyy"
RELLENO = PatternFill("solid", fgColor="1F3864")
BLANCO = Font(bold=True, color="FFFFFF")
NEGRITA = Font(bold=True)
LINEA = Border(top=Side(style="thin"))


def valor(dato):
    if pd.isna(dato):
        return None
    if isinstance(dato, pd.Timestamp):
        return dato.to_pydatetime()
    if hasattr(dato, "item"):
        return dato.item()
    return dato


def hojas(resultado):
    c = resultado.conciliados
    d = resultado.diferencias_monto
    b = resultado.solo_banco
    r = resultado.solo_registro
    u = resultado.duplicados
    regla = c["regla"].map({"operacion": "N° operación", "monto_fecha": "Monto y fecha"})
    return {
        "Conciliados": [
            ("Regla", regla, None),
            ("Fecha banco", c["fecha_banco"], FECHA),
            ("N° operación", c["operacion_banco"], None),
            ("Descripción", c["descripcion"], None),
            ("Monto banco", c["monto_banco"], MONEDA),
            ("Fecha libro", c["fecha_libro"], FECHA),
            ("Comprobante", c["comprobante"], None),
            ("Glosa", c["glosa"], None),
            ("Monto libro", c["monto_libro"], MONEDA),
        ],
        "Solo en banco": [
            ("Fecha", b["fecha"], FECHA),
            ("N° operación", b["operacion"], None),
            ("Descripción", b["descripcion"], None),
            ("Monto", b["monto"], MONEDA),
        ],
        "Solo en registro": [
            ("Fecha", r["fecha"], FECHA),
            ("Comprobante", r["comprobante"], None),
            ("Glosa", r["glosa"], None),
            ("Monto", r["monto"], MONEDA),
        ],
        "Diferencias de monto": [
            ("Fecha banco", d["fecha_banco"], FECHA),
            ("N° operación", d["operacion_banco"], None),
            ("Descripción", d["descripcion"], None),
            ("Monto banco", d["monto_banco"], MONEDA),
            ("Comprobante", d["comprobante"], None),
            ("Monto libro", d["monto_libro"], MONEDA),
            ("Diferencia", d["diferencia"], MONEDA),
        ],
        "Duplicados": [
            ("Fecha", u["fecha"], FECHA),
            ("Comprobante", u["comprobante"], None),
            ("Glosa", u["glosa"], None),
            ("N° operación", u["operacion"], None),
            ("Monto", u["monto"], MONEDA),
        ],
    }


def encabezado(celda):
    celda.font = BLANCO
    celda.fill = RELLENO
    celda.alignment = Alignment(horizontal="center", vertical="center")


def ajustar_anchos(hoja, minimo=10, maximo=50):
    for columna in hoja.columns:
        largo = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
        hoja.column_dimensions[get_column_letter(columna[0].column)].width = min(max(largo + 2, minimo), maximo)


def escribir_detalle(hoja, columnas):
    for j, (titulo, serie, formato) in enumerate(columnas, start=1):
        encabezado(hoja.cell(row=1, column=j, value=titulo))
        for i, dato in enumerate(serie, start=2):
            celda = hoja.cell(row=i, column=j, value=valor(dato))
            if formato:
                celda.number_format = formato
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    ajustar_anchos(hoja)
    for j, (_, _, formato) in enumerate(columnas, start=1):
        if formato == FECHA:
            hoja.column_dimensions[get_column_letter(j)].width = 14


def fila(hoja, n, concepto, monto, negrita=False, linea=False):
    a = hoja.cell(row=n, column=1, value=concepto)
    b = hoja.cell(row=n, column=2, value=monto)
    b.number_format = MONEDA
    if negrita:
        a.font = b.font = NEGRITA
    if linea:
        a.border = b.border = LINEA


def escribir_resumen(hoja, banco, saldos):
    desde, hasta = banco["fecha"].min(), banco["fecha"].max()
    hoja["A1"] = "Conciliación bancaria"
    hoja["A1"].font = Font(bold=True, size=14)
    hoja["A2"] = f"Periodo: {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"
    hoja["A3"] = f"Generado: {datetime.now():%d/%m/%Y %H:%M}"

    for j, titulo in enumerate(["Categoría", "Movimientos", "Monto"], start=1):
        encabezado(hoja.cell(row=5, column=j, value=titulo))
    categorias = [
        ("Conciliados", "Conciliados", "E"),
        ("Solo en el banco", "Solo en banco", "D"),
        ("Solo en el registro", "Solo en registro", "D"),
        ("Diferencias de monto", "Diferencias de monto", "G"),
        ("Duplicados", "Duplicados", "E"),
    ]
    for n, (nombre, hoja_detalle, columna) in enumerate(categorias, start=6):
        hoja.cell(row=n, column=1, value=nombre)
        hoja.cell(row=n, column=2, value=f"=COUNTA('{hoja_detalle}'!A:A)-1")
        hoja.cell(row=n, column=3, value=f"=SUM('{hoja_detalle}'!{columna}:{columna})").number_format = MONEDA

    for j, titulo in enumerate(["Cuadre de saldos", "Monto"], start=1):
        encabezado(hoja.cell(row=12, column=j, value=titulo))
    fila(hoja, 13, "Saldo según extracto bancario", saldos["saldo_banco"], negrita=True)
    fila(hoja, 14, "(+) Depósitos en tránsito", "=SUMIF('Solo en registro'!D:D,\">0\")")
    fila(hoja, 15, "(-) Pagos y cheques pendientes de cobro", "=SUMIF('Solo en registro'!D:D,\"<0\")")
    fila(hoja, 16, "Saldo bancario ajustado", "=SUM(B13:B15)", negrita=True, linea=True)
    fila(hoja, 18, "Saldo según libro bancos", saldos["saldo_libros"], negrita=True)
    fila(hoja, 19, "(+) Abonos del banco no registrados", "=SUMIF('Solo en banco'!D:D,\">0\")")
    fila(hoja, 20, "(-) Cargos del banco no registrados", "=SUMIF('Solo en banco'!D:D,\"<0\")")
    fila(hoja, 21, "(+/-) Corrección de diferencias de monto", "=SUM('Diferencias de monto'!G:G)")
    fila(hoja, 22, "(+/-) Reversión de asientos duplicados", "=-SUM(Duplicados!E:E)")
    fila(hoja, 23, "Saldo en libros ajustado", "=SUM(B18:B22)", negrita=True, linea=True)
    fila(hoja, 25, "Diferencia", "=B16-B23", negrita=True)

    hoja.column_dimensions["A"].width = 42
    hoja.column_dimensions["B"].width = 16
    hoja.column_dimensions["C"].width = 16


def generar_reporte(ruta_extracto, ruta_libro, ruta_salida):
    banco, saldo_inicial = leer_extracto(ruta_extracto)
    if saldo_inicial is None:
        raise ValueError("El extracto no indica el saldo inicial")
    libro = leer_libro(ruta_libro)
    resultado = conciliar(banco, libro)
    saldos = cuadre(resultado, banco, libro, saldo_inicial)

    libro_excel = Workbook()
    resumen = libro_excel.active
    resumen.title = "Resumen"
    for nombre, columnas in hojas(resultado).items():
        escribir_detalle(libro_excel.create_sheet(nombre), columnas)
    escribir_resumen(resumen, banco, saldos)

    ruta_salida = Path(ruta_salida)
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    libro_excel.save(ruta_salida)
    return resultado, saldos


def main():
    parser = argparse.ArgumentParser(description="Concilia un extracto bancario con el libro bancos.")
    parser.add_argument("--extracto", type=Path, default=Path("data/extracto_banco.csv"))
    parser.add_argument("--libro", type=Path, default=Path("data/libro_bancos.xlsx"))
    parser.add_argument("--salida", type=Path, default=Path("salida/conciliacion.xlsx"))
    args = parser.parse_args()

    try:
        resultado, saldos = generar_reporte(args.extracto, args.libro, args.salida)
    except PermissionError:
        sys.exit(f"No se pudo guardar {args.salida}. Si está abierto en Excel, ciérralo y vuelve a intentarlo.")
    except (FileNotFoundError, ValueError) as error:
        sys.exit(f"Error: {error}")

    print(f"Conciliados:          {len(resultado.conciliados):>3}")
    print(f"Solo en el banco:     {len(resultado.solo_banco):>3}")
    print(f"Solo en el registro:  {len(resultado.solo_registro):>3}")
    print(f"Diferencias de monto: {len(resultado.diferencias_monto):>3}")
    print(f"Duplicados:           {len(resultado.duplicados):>3}")
    print(f"Diferencia final:     {saldos['diferencia']:,.2f}")
    print(f"Reporte guardado en {args.salida}")


if __name__ == "__main__":
    main()
